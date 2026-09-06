import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.config import settings
from app.services.documents.security import DocumentIngestionError

PARSER_VERSION = "document-parser-v1"


@dataclass(frozen=True)
class SourceSpan:
    start: int
    end: int
    locator: dict[str, Any]


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    parser: str
    page_count: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)
    spans: list[SourceSpan] = field(default_factory=list)


def _clean(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = value.replace("\x00", "")
    value = re.sub(r"[\t ]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _bounded(text: str) -> str:
    cleaned = _clean(text)
    if len(cleaned) > settings.document_extracted_text_max_chars:
        raise DocumentIngestionError(
            "extracted_text_limit", "Extracted document text exceeds the configured limit"
        )
    return cleaned


def _decode_text(payload: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return payload.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise DocumentIngestionError("text_encoding", "Text encoding must be UTF-8 or GB18030")


def _parse_txt(payload: bytes) -> ParsedDocument:
    text, encoding = _decode_text(payload)
    return ParsedDocument(text=_bounded(text), parser="text", metadata={"encoding": encoding})


def _parse_csv(payload: bytes) -> ParsedDocument:
    decoded, encoding = _decode_text(payload)
    sample = decoded[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[str] = []
    cells = 0
    for row_number, row in enumerate(csv.reader(io.StringIO(decoded), dialect), start=1):
        if row_number > settings.document_workbook_max_rows:
            raise DocumentIngestionError("row_limit", "CSV exceeds the configured row limit")
        cells += len(row)
        if cells > settings.document_workbook_max_cells:
            raise DocumentIngestionError("cell_limit", "CSV exceeds the configured cell limit")
        safe = [str(cell)[:10_000] for cell in row]
        rows.append("\t".join(safe))
    return ParsedDocument(
        text=_bounded("\n".join(rows)),
        parser="csv",
        metadata={"encoding": encoding, "rows": len(rows), "cells": cells},
    )


def _parse_docx(payload: bytes) -> ParsedDocument:
    try:
        document = WordDocument(io.BytesIO(payload))
    except Exception as exc:
        raise DocumentIngestionError("docx_parse_failed", "DOCX content could not be parsed") from exc
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_number, table in enumerate(document.tables, start=1):
        lines.append(f"[Table {table_number}]")
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    text = _bounded("\n".join(lines))
    return ParsedDocument(
        text=text,
        parser="python-docx",
        metadata={"paragraphs": len(document.paragraphs), "tables": len(document.tables)},
    )


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ")[:10_000]


def _parse_xlsx(payload: bytes) -> ParsedDocument:
    try:
        workbook = load_workbook(
            io.BytesIO(payload), read_only=True, data_only=True, keep_links=False
        )
    except Exception as exc:
        raise DocumentIngestionError("xlsx_parse_failed", "XLSX content could not be parsed") from exc
    lines: list[str] = []
    rows = 0
    cells = 0
    try:
        for sheet in workbook.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                rows += 1
                if rows > settings.document_workbook_max_rows:
                    raise DocumentIngestionError(
                        "row_limit", "Workbook exceeds the configured row limit"
                    )
                cells += len(row)
                if cells > settings.document_workbook_max_cells:
                    raise DocumentIngestionError(
                        "cell_limit", "Workbook exceeds the configured cell limit"
                    )
                lines.append("\t".join(_cell_text(value) for value in row))
    finally:
        workbook.close()
    return ParsedDocument(
        text=_bounded("\n".join(lines)),
        parser="openpyxl",
        metadata={"sheets": len(workbook.sheetnames), "rows": rows, "cells": cells},
    )


def _parse_pdf(payload: bytes) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(payload), strict=False)
        if reader.is_encrypted:
            raise DocumentIngestionError("encrypted_pdf", "Encrypted PDF documents are not supported")
        parts: list[str] = []
        spans: list[SourceSpan] = []
        cursor = 0
        for page_number, page in enumerate(reader.pages, start=1):
            text = _clean(page.extract_text() or "")
            if not text:
                continue
            if parts:
                parts.append("\n\n")
                cursor += 2
            start = cursor
            parts.append(text)
            cursor += len(text)
            spans.append(SourceSpan(start=start, end=cursor, locator={"page": page_number}))
        joined = _bounded("".join(parts))
    except DocumentIngestionError:
        raise
    except Exception as exc:
        raise DocumentIngestionError("pdf_parse_failed", "PDF content could not be parsed") from exc
    if len(re.sub(r"\s", "", joined)) < 20:
        raise DocumentIngestionError(
            "ocr_required", "This PDF contains insufficient machine-readable text and requires OCR"
        )
    return ParsedDocument(
        text=joined,
        parser="pypdf",
        page_count=len(reader.pages),
        metadata={"pages_with_text": len(spans)},
        spans=spans,
    )


def parse_document(filename: str, payload: bytes) -> ParsedDocument:
    extension = Path(filename).suffix.lower()
    parser = {
        ".txt": _parse_txt,
        ".csv": _parse_csv,
        ".docx": _parse_docx,
        ".xlsx": _parse_xlsx,
        ".pdf": _parse_pdf,
    }.get(extension)
    if not parser:
        raise DocumentIngestionError("unsupported_type", "Unsupported document type")
    result = parser(payload)
    if not result.text:
        raise DocumentIngestionError("empty_document", "Document contains no extractable text")
    return result
