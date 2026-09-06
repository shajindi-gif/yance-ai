import io
import zipfile

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter
from weasyprint import HTML

from app.services.documents.intelligence import chunk_document, extract_company_intelligence
from app.services.documents.parsers import ParsedDocument, parse_document
from app.services.documents.security import (
    ClamAVScanner,
    DocumentIngestionError,
    verify_signature,
)

COMPANY_TEXT = """企业名称：上海衍策科技有限公司
所属行业：人工智能软件
注册地址：上海市徐汇区测试路 18 号
成立日期：2022年1月15日
员工人数：42 人
年度营业收入：1,200 万元
研发投入占营业收入 12.5%
融资阶段：A轮
高新技术企业认定：准备申报
科技型中小企业认定：已认定
专精特新企业认定：未认定
知识产权：软件著作权 3 项、发明专利 1 项
核心产品：企业政策智能平台、申报工作台
项目情况：企业知识智能化项目
当前需求：研发补贴、高企申报
"""


def _docx_bytes() -> bytes:
    document = WordDocument()
    document.add_heading("企业档案", level=1)
    document.add_paragraph(COMPANY_TEXT)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "企业档案"
    for line in COMPANY_TEXT.splitlines():
        label, _, value = line.partition("：")
        sheet.append([label, value])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _pdf_bytes() -> bytes:
    return HTML(
        string=(
            "<html><body><p>Company industry: artificial intelligence software.</p>"
            "<p>Annual revenue: 12 million yuan.</p></body></html>"
        )
    ).write_pdf()


@pytest.mark.parametrize(
    ("filename", "content_type", "payload", "expected_parser"),
    [
        ("profile.txt", "text/plain", COMPANY_TEXT.encode(), "text"),
        ("profile.csv", "text/csv", "字段,值\n行业,人工智能\n".encode(), "csv"),
        ("profile.pdf", "application/pdf", _pdf_bytes(), "pypdf"),
        (
            "profile.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _docx_bytes(),
            "python-docx",
        ),
        (
            "profile.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            _xlsx_bytes(),
            "openpyxl",
        ),
    ],
)
def test_real_document_parsers(filename, content_type, payload, expected_parser):
    verify_signature(filename, content_type, payload)
    parsed = parse_document(filename, payload)
    assert parsed.parser == expected_parser
    assert parsed.text


def test_image_only_pdf_is_explicitly_routed_to_ocr():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = io.BytesIO()
    writer.write(stream)
    payload = stream.getvalue()
    verify_signature("scan.pdf", "application/pdf", payload)
    with pytest.raises(DocumentIngestionError, match="requires OCR") as error:
        parse_document("scan.pdf", payload)
    assert error.value.code == "ocr_required"


def test_signature_mismatch_and_binary_text_are_rejected():
    with pytest.raises(DocumentIngestionError) as mismatch:
        verify_signature("fake.pdf", "application/pdf", b"not-a-pdf")
    assert mismatch.value.code == "signature_mismatch"

    with pytest.raises(DocumentIngestionError) as binary:
        verify_signature("fake.txt", "text/plain", b"hello\x00world")
    assert binary.value.code == "binary_text"


def test_office_structure_must_match_extension():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
    with pytest.raises(DocumentIngestionError) as mismatch:
        verify_signature(
            "fake.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            stream.getvalue(),
        )
    assert mismatch.value.code == "office_structure_invalid"


def test_company_intelligence_is_structured_and_source_grounded():
    parsed = ParsedDocument(text=COMPANY_TEXT, parser="text")
    candidates = {item.field_key: item for item in extract_company_intelligence(parsed)}
    assert candidates["employee_count"].value == 42
    assert candidates["annual_revenue"].value == 12_000_000
    assert candidates["rd_expense_ratio"].value == 12.5
    assert candidates["founded_date"].value == "2022-01-15"
    assert candidates["core_products"].value == ["企业政策智能平台", "申报工作台"]
    assert "年度营业收入" in candidates["annual_revenue"].evidence_excerpt
    assert candidates["annual_revenue"].source_locator["character_start"] >= 0
    assert len({item.candidate_hash for item in candidates.values()}) == len(candidates)


def test_chunks_are_bounded_and_overlap(monkeypatch):
    monkeypatch.setattr("app.services.documents.intelligence.settings.document_chunk_size_chars", 80)
    monkeypatch.setattr("app.services.documents.intelligence.settings.document_chunk_overlap_chars", 10)
    chunks = chunk_document(ParsedDocument(text=COMPANY_TEXT * 3, parser="text"))
    assert len(chunks) > 1
    assert all(len(item["content"]) <= 80 for item in chunks)
    assert chunks[1]["character_start"] < chunks[0]["character_end"]


def test_disabled_malware_scanner_is_explicit_not_fake_clean(monkeypatch):
    monkeypatch.setattr("app.services.documents.security.settings.malware_scan_mode", "disabled")
    result = ClamAVScanner().scan(b"safe sample")
    assert result.status == "skipped"
    assert result.engine == "disabled"


def test_required_malware_scanner_fails_closed_when_unavailable(monkeypatch):
    monkeypatch.setattr("app.services.documents.security.settings.malware_scan_mode", "required")
    monkeypatch.setattr("app.services.documents.security.settings.clamd_host", "scanner.invalid")
    monkeypatch.setattr(
        "app.services.documents.security.socket.create_connection",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("unavailable")),
    )
    with pytest.raises(DocumentIngestionError) as error:
        ClamAVScanner().scan(b"sample")
    assert error.value.code == "malware_scanner_unavailable"
    assert error.value.retryable is True
